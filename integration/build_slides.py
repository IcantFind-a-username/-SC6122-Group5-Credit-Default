"""Build editable 16-slide presentation and native notes from audited artifacts.

Commit before execution. Run: PYTHONDONTWRITEBYTECODE=1 python -m integration.build_slides
Charts remain editable PowerPoint charts. PowerPoint or a supplied LibreOffice runtime renders the actual PPTX to PDF.
"""
import argparse
from dataclasses import dataclass
from pathlib import Path
import json
import subprocess
import tempfile

import fitz
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

from integration.artifacts import file_hash, write_json
from integration.audit import read_csv
from integration.presentation_evidence import rf_interpretation

ROOT = Path(__file__).resolve().parents[1]
NAVY, TEAL, ORANGE = "183447", "187B78", "C8784A"
LIGHT, GRAY, WHITE = "E1E8E8", "526575", "FFFFFF"
PAPER = "FAFAF7"


@dataclass
class Slide:
    title: str
    role: str
    seconds: int
    bullets: list[str]
    headline: str
    detail: str
    source: str
    script: str
    chart: dict | None = None
    transition: str = ""


def textbox(slide, x, y, w, h, text, size=22, color=NAVY, bold=False):
    """Presentation layout helper; does not rasterize editable text."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.03)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(7)
    return box


def remove_shadow(shape):
    """Disable both local and theme effects for consistent Office/PDF rendering."""
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    for reference in shape._element.xpath("p:style/a:effectRef"):
        reference.set("idx", "0")


def panel(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    remove_shadow(shape)


def rule(slide, x1, y1, x2, y2, color=LIGHT, width=1):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = RGBColor.from_string(color)
    line.line.width = Pt(width)
    remove_shadow(line)


def editable_chart(slide, spec):
    """Native chart with embedded workbook, direct values and a zero baseline."""
    data = CategoryChartData()
    data.categories = spec["categories"]
    for name, values in spec["series"]:
        data.add_series(name, values)
    chart_type = XL_CHART_TYPE.BAR_CLUSTERED if spec.get("horizontal") else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(chart_type,
        Inches(.65), Inches(2.45), Inches(7.3), Inches(3.75), data).chart
    chart.has_title = False
    chart.has_legend = len(spec["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(15)
    chart.category_axis.tick_labels.font.size = Pt(16)
    chart.value_axis.tick_labels.font.size = Pt(14)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = spec.get("maximum", 1)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(LIGHT)
    chart.category_axis.format.line.color.rgb = RGBColor.from_string(LIGHT)
    chart.value_axis.format.line.fill.background()
    chart.value_axis.tick_labels.number_format = spec.get("format", "0.0")
    plot = chart.plots[0]
    plot.gap_width = 90
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(17)
    plot.data_labels.number_format = spec.get("format", "0.000")
    palette = ("899CA5", ORANGE if spec.get("format") == "0" else TEAL)
    for index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(palette[index % 2])
        series.format.line.fill.background()
        if len(spec["series"]) == 1:
            for point_index, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(TEAL if spec.get("uniform_color") else palette[point_index % 2])
                point.format.line.fill.background()


def make_slides(frame):
    def row(family, model):
        selected = frame[(frame.Family == family) & (frame.Partition == "test") & (frame.Model == model)]
        if len(selected) != 1:
            raise ValueError(f"Expected one audited row for {family}/{model}")
        return selected.iloc[0]

    lr = row("Logistic Regression", "Tuned LR supplement")
    lb = row("Logistic Regression", "Baseline LR supplement")
    dt = row("Decision Tree", "Tuned Decision Tree")
    db = row("Decision Tree", "Baseline Decision Tree")
    rf = row("RF", "Tuned RF")
    rb = row("RF", "Baseline RF")
    xg = row("XGBoost", "Tuned XGBoost")
    xb = row("XGBoost", "Baseline XGBoost")
    xc = row("XGBoost", "Tuned XGBoost / cost ratio 5")
    rf_parameters = json.loads((ROOT / "results/rf/protocol_frozen.json").read_text())["Selected_parameters"]
    xg_parameters = json.loads((ROOT / "results/xgboost/protocol_frozen.json").read_text())["Selected_parameters"]
    xg_interval = read_csv(ROOT / "results/xgboost/bootstrap_intervals.csv").set_index("Quantity").loc["AP difference: tuned - baseline"]
    importance, examples = rf_interpretation()
    fn, fp = examples.loc["FN"], examples.loc["FP"]
    top_importance = importance.head(5).iloc[::-1]
    source = "Source: results/final/model_comparison.csv | Historical test n=6,000; default=1"
    note = "LR* = post-hoc reproducibility supplement after historic test results existed."
    comparison = {"categories": ["LR*", "Tree", "RF", "XGBoost"],
                  "series": [("Baseline", [lb.AP, db.AP, rb.AP, xb.AP]),
                             ("CV-selected", [lr.AP, dt.AP, rf.AP, xg.AP])],
                  "maximum": .65, "format": "0.000", "horizontal": True}
    slides = [
        Slide("Credit Card Default Prediction", "Part 1", 45,
          ["How well can we identify default risk?", "What changes when missed defaults cost more?", "Four models. One auditable comparison."],
          "Risk, interpretation and decision costs", "NTU SC6122\nEmerging Topics in FinTech\nGroup 5", "UCI Default of Credit Card Clients; group experiment artifacts.",
          "Our question is how to identify default risk while balancing predictive ranking, interpretability, and the cost of taking action. We compare logistic regression, a decision tree, random forest, and XGBoost on shared historical records. A model score ranks accounts, but a threshold decides which accounts receive attention. Those are different questions, and our results will show why the distinction matters. I am Lei Peng, and I will begin with the data and logistic regression. Zhang Hanyu will cover the decision tree, Zhou Xinzhe the random forest, and Xu Yiqun XGBoost and decision costs."),
        Slide("How we split and use the data", "Part 1", 70,
          ["30,000 clients • 23 predictors • 22.12% default", "Categories learned within each training fold", "Validation selects policies; historical test reports outcomes"],
          "18,000 → 6,000 → 6,000", "FIT: five CV folds, seed 42\nVALIDATION: threshold selection\nTEST: fixed historical holdout", "Source: data/splits; protocol_frozen.json; results/final/audit.json",
          "The UCI dataset contains thirty thousand Taiwan credit card clients and twenty-three predictors. About twenty-two percent default, so an always-negative classifier can appear accurate while detecting nobody. We retain the supplied stratified split: eighteen thousand fitting rows, six thousand validation rows, and six thousand historical test rows. Five shared folds inside fitting data select configurations by average precision. Each fold learns its own category encoding, and logistic regression also learns scaling there. Validation selects thresholds without a later refit. The audit confirms identical memberships across the reproducible models. This random historical holdout is not temporal or external validation."),
        Slide("Logistic regression: the linear reference", "Part 1", 65,
          ["L2 model + fold-trained scaling + one-hot categories", "12 declared settings; CV selected C=0.1, unweighted", "Original LR notebook was empty; original process unverified"],
          f"AP {lr.AP:.4f}", f"Baseline AP {lb.AP:.4f}\nNo ranking improvement from selection\nLR* is a separate post-hoc supplement", source + " | " + note,
          f"Logistic regression offers a linear reference whose coefficients describe regularized associations. We used fold-trained scaling and categorical encoding, then compared twelve declared combinations of regularization strength and class weighting. Cross-validation selected C equal to zero point one without class weighting. The test AP is {lr.AP:.4f}, versus {lb.AP:.4f} for the baseline, so selection did not improve ranking here. Crucially, the original logistic notebook was empty, and its original procedure cannot be verified. We preserve those artifacts and label our reproducible run as a post-hoc supplement after earlier test results existed. Part two now examines how tree regularization changes performance."),
        Slide("Decision tree: controlling overfitting", "Part 2", 50,
          ["21 configurations: depth × minimum leaf size", "Unrestricted baseline fits training folds almost perfectly", "Selected minimum leaf size: 100 observations"],
          "Test average precision", "A large baseline gap signals overfitting.\nCV selection improves this tree's ranking.", source,
          "A decision tree can create very detailed splits, so it can easily overfit.\nIn our baseline model, the tree fits the training folds almost perfectly, but its cross-validation performance is much lower.\nWe therefore tested 21 combinations of maximum depth and minimum leaf size. The selected model requires at least 100 observations in each leaf, while keeping no explicit depth limit.\nOn the test set, average precision improves from 0.284 to 0.522.\nSo the key point is that controlling leaf size greatly improves the decision tree’s held-out performance on this split.\nNext, I’ll show one of the main rules learned by this model.",
          {"categories": ["Baseline", "Selected"], "series": [("AP", [db.AP, dt.AP])], "maximum": .65}),
        Slide("Reading the tree’s first decision", "Part 2", 55,
          ["Root rule: recent repayment status PAY_0 ≤ 1.5", "Selected tree: depth 18, with 134 leaves", "Node proportions describe this fitting sample"],
          "16.51%  |  70.13%", "Fit-row default proportions\nLeft branch | Right branch\nAssociation, not a causal effect", "Source: saved decision-tree model; integration audit/report evidence.",
          "The selected tree starts with recent repayment status, PAY zero.\nThe root split is at 1.5. In the fitting sample, the default rate is about 16.5% on the left branch, compared with about 70.1% on the right branch.\nThis shows that recent repayment status is one of the strongest predictive signals used by the tree.\nHowever, this is only an association in our fitted data. It does not mean repayment status directly causes default.\nThe full tree is still quite large, with depth 18 and 134 leaves, so this rule is just one representative example rather than a complete explanation of the model.\nNext, we compare the tree with the other models using the same ranking metric."),
        Slide("Ranking performance across four models", "Part 2", 75,
          ["Decision tree gains most on this split", "RF and XGBoost have the highest selected AP", "CV selection precedes test comparison"],
          "Test average precision", "Baseline and selected models are both retained.\nSmall ensemble differences do not establish superiority.", source + " | " + note,
          "Finally, we compare the effect of tuning across all four model families using test average precision.\nThe most noticeable result is the decision tree. Its AP increases substantially from 0.284 to 0.522 after tuning. In comparison, logistic regression stays around 0.497, and random forest improves only moderately. XGBoost changes very little.\nIn this comparison, the decision tree benefits the most from hyperparameter tuning, which is consistent with what we saw earlier: controlling tree complexity is especially important for reducing overfitting.\nIn terms of final ranking performance, the selected XGBoost and random forest models achieve AP values of 0.552 and 0.551 respectively. These values are extremely close, so we should not claim that one is clearly superior.\nAlso, model selection was based on cross-validation before evaluating the test set, so we do not re-select a model simply because a baseline happens to perform slightly better on the test data.\nLR remains a separately labelled post-hoc supplement. Zhou Xinzhe will now explain random-forest tuning, feature importance and error cases.", comparison),
        Slide("Random forest: averaging and regularization", "Part 3", 70,
          ["Baseline + 23 seeded candidate settings", f"Selected: {rf_parameters['n_estimators']} trees, depth {rf_parameters['max_depth']}, minimum leaf {rf_parameters['min_samples_leaf']}", f"Positive-class weight {rf_parameters['class_weight']['1']}; feature fraction {rf_parameters['max_features']}"],
          "Test average precision", "Averaging reduces dependence on one tree.\nWeights also change the score distribution.", source + " | results/rf/protocol_frozen.json",
          f"Random forest averages many trees, reducing dependence on any single partition. We evaluated a baseline and twenty-three seeded candidate configurations using the same five fitting folds and average-precision objective. The selected configuration uses five hundred trees, maximum depth eight, a minimum leaf size of two, and half the features at each split. It also weights the positive class three times as heavily. Test AP changes from {rb.AP:.4f} to {rf.AP:.4f}. Class weighting affects the learned score distribution, so we should not assume these scores are calibrated probabilities. We next examine which features the fitted forest relies on and where it makes errors.",
          {"categories": ["Baseline", "Selected"], "series": [("AP", [rb.AP, rf.AP])], "maximum": .65}),
        Slide("Random forest: which features matter?", "Part 3", 60,
          ["PAY_0 has the largest validation AP decrease", "Permutation measures predictive reliance", "Correlated variables can substitute for each other"],
          "Mean decrease in validation AP", "Five saved shuffles per feature.\nPredictive association, not causality.", "Source: results/rf/validation_importance.csv | Five validation permutations; frozen model.",
          f"We measure importance by shuffling one feature on validation data and observing the fall in average precision. The forest relies most strongly on recent repayment status, PAY zero: its mean AP decrease is {importance.iloc[0].Mean_AP_decrease:.4f}. Earlier repayment status, PAY two, follows at {importance.iloc[1].Mean_AP_decrease:.4f}. These values come from five saved permutations of each feature. They describe the fitted forest's predictive reliance, not a causal effect or a feature's isolated contribution. Correlated predictors can substitute for one another, and shuffling may create combinations that do not occur naturally. The chart shows mean decreases; variation across shuffles is recorded in the report. Even an influential variable cannot explain every customer's outcome.",
          {"categories": top_importance.Feature.tolist(), "series": [("Mean AP decrease", top_importance.Mean_AP_decrease.tolist())],
           "maximum": .24, "format": "0.000", "horizontal": True, "uniform_color": True}),
        Slide("Random forest: two illustrative errors", "Part 3", 50,
          [f"At 0.5: {int(rf.FN):,} missed defaults; {int(rf.FP):,} false alarms",
           f"Missed default · row {int(fn.row_id)}\nScore {fn.Tuned_probability:.4f}; actual default = 1",
           f"False alarm · row {int(fp.row_id)}\nScore {fp.Tuned_probability:.4f}; actual default = 0"],
          "Why the predictions can fail", f"Missed default: PAY_0 = {int(fn.PAY_0)}\nFalse alarm: PAY_0 = {int(fp.PAY_0)}\n\nIllustrative extreme errors, not typical clients.",
          "Source: results/rf/misclassification_cases_05.csv | Tuned RF; score ≥ 0.5 predicts default.",
          f"At the fixed threshold of zero point five, the selected forest misses {int(rf.FN)} defaults and falsely flags {int(rf.FP)} non-defaults. Consider two examples already present in the saved error export. Row {int(fn.row_id)} actually defaults but receives a score of {fn.Tuned_probability:.4f}; its recent repayment code is {int(fn.PAY_0)}. Row {int(fp.row_id)} does not default but scores {fp.Tuned_probability:.4f}; its recent repayment code is {int(fp.PAY_0)}. These examples show that repayment history is informative but not deterministic. They are the most confident errors from the existing export, not representative average customers. We do not change the model after inspecting them. Xu Yiqun will now examine the separate trade-off created by changing a validation-selected threshold."),
        Slide("XGBoost: no gain in test ranking", "Part 4", 60,
          ["24 configurations; select by CV AP", f"Selected: {xg_parameters['n_estimators']} trees, depth {xg_parameters['max_depth']}, rate {xg_parameters['learning_rate']}", f"Selected ROC-AUC: {xg['ROC-AUC']:.4f}"],
          "Test average precision", f"Tuned − baseline AP: {xg_interval.Estimate:.4f}\nConditional 95% interval: [{xg_interval.CI_low:.4f}, {xg_interval.CI_high:.4f}]", source + " | results/xgboost/bootstrap_intervals.csv",
          f"XGBoost builds trees sequentially with regularization. We retained the baseline and searched twenty-three additional seeded configurations using the shared CV objective. The selected configuration has test AP {xg.AP:.4f}, compared with {xb.AP:.4f} for the baseline. The paired bootstrap interval for the difference spans zero, so we do not claim a ranking improvement. This bootstrap resamples the fixed predictions and does not include model retraining or selection uncertainty. The useful next question is separate: given the selected model's scores, how does a validation-chosen action threshold change missed defaults, false positives, and the stated hypothetical cost?",
          {"categories": ["Baseline", "Selected"], "series": [("AP", [xb.AP, xg.AP])], "maximum": .65}),
        Slide("XGBoost: fewer misses, more reviews", "Part 4", 80,
          [f"Validation r=5 cutoff: {xc.Threshold:.4f}", f"Recall {xg.Recall:.1%} → {xc.Recall:.1%}; alerts {xc.Alert_rate:.1%}", f"Cost {int(xg.Cost_5):,} → {int(xc.Cost_5):,}: {(1-xc.Cost_5/xg.Cost_5):.1%} reduction"],
          "Test confusion trade-off", "Same ranking scores, different action policy.\nReview costs are hypothetical; capacity is unmeasured.", source + " | FP + 5×FN; threshold selected only on validation.",
          f"At a cost ratio of five, validation selects XGBoost's threshold at {xc.Threshold:.4f}. Applying it unchanged to the historical test set reduces false negatives from {int(xg.FN)} to {int(xc.FN)}, while false positives rise from {int(xg.FP)} to {int(xc.FP)}. Recall reaches {xc.Recall:.1%}, but {xc.Alert_rate:.1%} of clients are flagged. The total hypothetical cost falls from {int(xg.Cost_5)} to {int(xc.Cost_5)}, about {(1-xc.Cost_5/xg.Cost_5):.1%}. This is not an improvement in AP or ROC-AUC; the scores have not changed. The policy may be impractical if review capacity is limited. Real exposure, loss severity, and intervention benefits are unavailable, so these cost units are a sensitivity analysis.",
          {"categories": ["False positives", "False negatives"], "series": [("Threshold 0.5", [xg.FP, xg.FN]), ("Validation r=5", [xc.FP, xc.FN])], "maximum": 2700, "format": "0"}),
        Slide("What the results mean for credit review", "Part 4", 40,
          ["Ranking, explanations, and workload need separate evidence", "No clear ensemble winner from these small differences", "Next: fresh temporal data, calibration, subgroup errors"],
          "Discussion", "Set realistic costs and review capacity.\nValidate the policy on fresh data.", "Historical educational study; no causal or deployment-readiness claim.",
          "Ranking and action policy answer different questions. Tree regularization improves its baseline, while the selected ensembles have similar ranking estimates. Lower thresholds reduce the stated cost by creating more reviews. Operational recommendations require realistic costs and capacity constraints. Historical data, duplicates, random splitting, and logistic provenance remain limitations. A new study should use fresh temporal or external data and examine calibration and subgroup errors. Thank you. We welcome your questions."),
        Slide("Backup A • AP, ROC-AUC and threshold metrics", "Q&A", 0,
          ["AP = sum of recall increments × precision", "ROC-AUC measures positive/negative score ordering", "Precision, recall, F1 and accuracy require a threshold"],
          "Two different definitions", "Saved continuous scores supply ranking metrics.\nChanging the cutoff changes decisions only.", "Source: scikit-learn average_precision_score documentation; part4/evaluation.py",
          "If asked about the ranking metric, explain that average precision weights each precision value by the associated increase in recall. It does not linearly interpolate the precision-recall curve with trapezoids, so similarly named PR-AUC values need not agree. ROC-AUC instead summarizes positive-versus-negative score ordering. Neither metric changes when we apply a different operating threshold to the same saved scores. Precision and recall do change because the decision set changes. The historical logistic CSV uses a PR-AUC label whose computation cannot be verified, which is one reason it is excluded from our verified ranking table."),
        Slide("Backup B • Cost ratios and fixed-prediction uncertainty", "Q&A", 0,
          ["Predeclared ratios: r = 1, 3, 5, 10", "Validation minimum cost; ties prefer largest cutoff", "1,000 stratified paired bootstrap resamples"],
          "What the intervals cover", "Fixed models / thresholds / class counts\nExcludes fitting / selection uncertainty\nExcludes population / prevalence shifts", "Source: frozen protocols; integration/uncertainty.py; bootstrap_intervals.csv",
          "The cost experiment declares four ratios rather than searching for a convenient ratio after viewing test results. At each ratio, the validation table includes all distinct score-based decision sets, with tied scores moving together. Equal costs select the largest threshold and therefore fewer alerts. The paired bootstrap samples the same test-row indices for each compared prediction vector, stratified by outcome. Its intervals are conditional on frozen models, thresholds, and class counts. They do not include uncertainty from refitting, hyperparameter search, threshold search, or a shift to another population. Real operational costs would require information that this dataset does not provide."),
        Slide("Backup C • Data provenance and evidence boundaries", "Q&A", 0,
          ["UCI cleaning reproduced exactly from downloaded source", "EDUCATION 0/5/6 → 4; MARRIAGE 0 → 3", "Retained duplicate vectors can span partitions"],
          "The limits of historical evidence", "Original LR process remains unverified.\nLR* source was committed before its one run.\nNo new unseen-test claim is made.", "Sources: results/final/source_audit.json; integration/LOGISTIC_PROVENANCE.md",
          "The integration audit redownloaded the official source and reproduced the committed cleaned dataset in values, types, order, and schema. Undocumented education categories are consolidated, and marriage category zero is mapped to other. The original download bytes were not preserved, so the new download and its hash are documented separately. Duplicate feature vectors remain under the established protocol and may create dependence across partitions. For logistic regression, the original notebook history contained an empty model notebook. The supplement's code and search declaration were committed before execution, but previous team test results already existed. This makes the supplement reproducible without making it independent external evidence."),
        Slide("Backup D • Sources, responsibilities and reproducibility", "Q&A", 0,
          ["UCI: doi.org/10.24432/C55S3H", "scikit-learn: average_precision_score documentation", "XGBoost paper: doi.org/10.1145/2939672.2939785"],
          "Member contributions", "Equal contribution: 25% per member", "Sources: Group 5 repository, frozen protocols, row-level predictions and integration audit.",
          "The numerical source for this editable presentation is the audited final comparison CSV, with model-specific protocols and bootstrap artifacts supplying supporting context. The UCI dataset, scikit-learn average-precision documentation, and the XGBoost paper are the core external references. The role plan assigns three minutes each to data and logistic regression, decision trees and ranking, random forest with feature importance and error cases, and XGBoost with threshold trade-offs. Lei Peng covers part one, Zhang Hanyu part two, Zhou Xinzhe part three, and Xu Yiqun part four. The group has confirmed equal contributions of twenty-five percent each. Data-split and metric agreement, final integration, checking and rehearsal are shared by all four members. Integration and the logistic supplement were prepared with AI assistance for member review. Reproducing artifacts is distinct from retraining and selecting new models."),
    ]
    transitions = [
        "First, let us see how the shared data supports a fair comparison.",
        "With the data roles fixed, we can introduce the logistic reference and its caveat.",
        "I will now hand over to Zhang Hanyu to explain decision-tree regularization.",
        "Next, we inspect one rule to understand the tree's predictive pattern.",
        "We can now place the tree beside the other models using a common ranking metric.",
        "I will now hand over to Zhou Xinzhe for random forest, feature importance, and error cases.",
        "We next examine which features the forest uses to rank risk.",
        "We now look at two saved errors to understand the limits of these patterns.",
        "I will now hand over to Xu Yiqun for XGBoost and the cost of acting on its scores.",
        "The next slide holds scores fixed and changes only the validation-selected threshold.",
        "This trade-off leads to our main conclusion about ranking, costs, and workload.",
        "We welcome your questions; the following slides provide supporting details.",
        "If useful, we can next explain the cost rule and the scope of its uncertainty.",
        "We can also explain where the source data and evidence boundaries come from.",
        "The last backup lists the references and each member’s contribution.",
        "We can return to the main conclusion or discuss a specific model or policy.",
    ]
    for slide, transition in zip(slides, transitions, strict=True):
        slide.transition = transition
    return slides


def cover(slide, spec, team):
    """A typographic cover with the complete team, without rehearsal timings."""
    panel(slide, 0, 0, 13.333, 7.5, PAPER)
    panel(slide, 8.1, 0, 5.233, 7.5, NAVY)
    textbox(slide, .7, .6, 7.0, .45, "NTU  /  SC6122  /  GROUP 5", 14, TEAL, True)
    textbox(slide, .7, 1.72, 7.0, 1.9, "Credit Card\nDefault Prediction", 44, NAVY, True)
    rule(slide, .73, 4.0, 2.0, 4.0, TEAL, 2)
    textbox(slide, .7, 4.3, 6.6, 1.2, spec.headline, 25, GRAY)
    textbox(slide, .7, 6.4, 6.8, .5, "Emerging Topics in FinTech", 16, GRAY)
    textbox(slide, 8.55, .65, 4.2, .5, "THE TEAM", 13, "A6C8CB", True)
    for index, member in enumerate(team):
        y = 1.55 + index * 1.24
        textbox(slide, 8.55, y, 4.15, .4, member['name'], 23, WHITE, True)
        textbox(slide, 8.55, y + .45, 4.15, .4,
                f"{member['student_id']}  ·  {member['role']}", 14, "CBDADC")
        if index < 3:
            rule(slide, 8.58, y + .98, 12.65, y + .98, "385365", .7)
    textbox(slide, 8.55, 6.75, 4.1, .35, "Contribution: 25% each", 13, "CBDADC")


def content_slide(slide, spec, index, member, team):
    """Open layouts: charts, a data table, a tree rule, or a short text column."""
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor.from_string(WHITE)
    section = spec.role if spec.seconds else "Q&A backup"
    textbox(slide, .7, .42, 11.9, .4,
            f"{section.upper()}  /  {member['name']}  /  {member['scope']}", 13, TEAL, True)
    textbox(slide, .7, 1.0, 12.0, .86, spec.title, 31, NAVY, True)
    rule(slide, .73, 1.91, 12.58, 1.91)
    if spec.chart:
        textbox(slide, .76, 2.1, 7.1, .35, spec.headline, 16, TEAL, True)
        editable_chart(slide, spec.chart)
        rule(slide, 8.27, 2.35, 8.27, 6.15)
        for j, bullet in enumerate(spec.bullets):
            textbox(slide, 8.65, 2.22 + j * 1.01, 3.98, .89, bullet, 20, NAVY, j == 0)
        textbox(slide, 8.65, 5.43, 3.98, .94, spec.detail, 15, GRAY)
    else:
        for j, bullet in enumerate(spec.bullets):
            y = 2.37 + j * 1.22
            textbox(slide, .72, y, 5.45, 1.03, bullet, 23)
            if j < 2:
                rule(slide, .75, y + 1.08, 5.94, y + 1.08)
        rule(slide, 6.42, 2.4, 6.42, 6.15)
        if index == 2:
            textbox(slide, 6.85, 2.3, 5.4, .45, "Partition and purpose", 21, TEAL, True)
            for j, (label, count, purpose) in enumerate([
                ("Fitting", "18,000", "Five-fold CV and final fitting"),
                ("Validation", "6,000", "Select decision thresholds"),
                ("Test", "6,000", "Evaluate frozen decisions")]):
                y = 3.05 + j * 1.05
                textbox(slide, 6.85, y, 3.6, .4, label, 21, NAVY, True)
                textbox(slide, 10.68, y, 1.7, .4, count, 21, TEAL)
                textbox(slide, 6.85, y + .45, 5.3, .35, purpose, 16, GRAY)
                rule(slide, 6.88, y + .9, 12.5, y + .9)
        elif index == 5:
            textbox(slide, 6.86, 2.4, 5.5, .5, "Recent repayment status", 21, TEAL, True)
            textbox(slide, 8.35, 3.25, 3.2, .55, "PAY_0 ≤ 1.5", 25, NAVY, True)
            rule(slide, 9.7, 3.95, 8.05, 4.62, TEAL, 1.5)
            rule(slide, 9.7, 3.95, 11.25, 4.62, TEAL, 1.5)
            textbox(slide, 7.23, 4.8, 2.7, .45, "Yes · 16.51%", 25, TEAL)
            textbox(slide, 10.25, 4.8, 2.3, .45, "No · 70.13%", 25, ORANGE)
            textbox(slide, 6.9, 5.62, 5.65, .75,
                    "Default proportion in each fitting-sample branch.\nThese are associations, not causal effects.", 16, GRAY)
        elif index == 16:
            textbox(slide, 6.85, 2.3, 5.6, .5, "Member contributions · 25% each", 21, TEAL, True)
            for j, entry in enumerate(team):
                y = 3.0 + j * .8
                textbox(slide, 6.85, y, 5.7, .4,
                        f"{entry['name']}  ·  {entry['student_id']}", 18, NAVY, True)
                textbox(slide, 6.85, y + .38, 5.7, .35,
                        f"{entry['role']} — {entry['scope']}", 15, GRAY)
        else:
            textbox(slide, 6.88, 2.55, 5.5, 1.25, spec.headline, 31, TEAL, True)
            rule(slide, 6.92, 4.03, 8.1, 4.03, TEAL, 2)
            textbox(slide, 6.88, 4.4, 5.55, 1.75, spec.detail, 21, GRAY)
    source = spec.source.split(" | ")[0]
    textbox(slide, .72, 6.64, 11.65, .38, source, 10, GRAY)
    rule(slide, .73, 7.09, 12.59, 7.09)
    textbox(slide, .72, 7.18, 10.9, .27, "SC6122  ·  GROUP 5", 10, GRAY)
    textbox(slide, 12.13, 7.13, .5, .32, f"{index:02}", 12, TEAL, True)


def run(soffice=None):
    output = ROOT / "submission"
    output.mkdir(exist_ok=True)
    frame = read_csv(ROOT / "results/final/model_comparison.csv")
    slides = make_slides(frame)
    assert len(slides) == 16 and sum(s.seconds for s in slides) == 720
    for role in ["Part 1", "Part 2", "Part 3", "Part 4"]:
        assert sum(s.seconds for s in slides if s.role == role) == 180
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    team = json.loads((output / "team.json").read_text())
    by_role = {member['role']: member for member in team}
    backup_leads = {13: 'Part 2', 14: 'Part 4', 15: 'Part 1', 16: 'Part 3'}
    notes = ["# Group 5 — English speaker notes", "12-minute talk + 3-minute Q&A. Each member: 180 seconds.",
             "Timing is a rehearsal guide and appears only in notes, not on audience slides.",
             "All scripts below are also embedded in the PPTX speaker notes for Presenter View.",
             "## Speaker assignments"]
    assignment_table = ["| Speaker / student ID | Section | Slides | Contribution |", "|---|---|---|---|"]
    assignment_table += [f"| {m['name']} / {m['student_id']} | {m['role']}: {m['contribution_scope']} | {m['slides']} | {m['share_percent']}% |" for m in team]
    notes.append("\n".join(assignment_table))
    notes += ["Shared by all four members: agree splits, preprocessing and metrics; write individual report/slides; jointly integrate, check results and rehearse. Q15–Q16 audit questions are shared preparation.",
              "Q&A backup leads: slide 13 Zhang Hanyu; slide 14 Xu Yiqun; slide 15 Lei Peng; slide 16 Zhou Xinzhe. Other members support questions in their model area.",
              "Integration and LR supplement prepared with AI assistance for member review."]
    speaker_map = {}
    for index, spec in enumerate(slides, 1):
        member = by_role[spec.role if spec.seconds else backup_leads[index]]
        speaker_map[str(index)] = member['name']
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if index == 1:
            cover(slide, spec, team)
        else:
            content_slide(slide, spec, index, member, team)
        assignment = f"{member['name']} ({member['student_id']}) | {member['role']}: {member['scope']} | Contribution {member['share_percent']}%"
        timing = f"Target: {spec.seconds} seconds" if spec.seconds else "Q&A backup — use when relevant"
        native = f"Slide {index}: {spec.title}\n{assignment}\n{timing}\n\n{spec.script}\n\nTransition: {spec.transition}\n\n{spec.source}"
        slide.notes_slide.notes_text_frame.text = native
        notes += [f"## Slide {index}: {spec.title}", f"**{assignment}**", timing, spec.script,
                  f"**Transition:** {spec.transition}", spec.source]
    pptx_path = output / "Group5_Presentation.pptx"
    pdf_path = output / "Group5_Presentation.pdf"
    deck.save(str(pptx_path))
    notes_path = output / "Group5_Speaker_Notes.md"
    notes_path.write_text("\n\n".join(notes) + "\n", encoding="utf-8")
    export_script = '''on run argv
tell application "Microsoft PowerPoint"
open (POSIX file (item 1 of argv))
set d to active presentation
save d in (POSIX file (item 2 of argv)) as save as PDF
close d saving no
end tell
end run'''
    if soffice is None:
        subprocess.run(["osascript", "-e", export_script, str(pptx_path), str(pdf_path)], check=True, timeout=180)
        renderer = "Microsoft PowerPoint native PDF export"
        renderer_version = None
    else:
        with tempfile.TemporaryDirectory(prefix="group5-libreoffice-") as profile:
            subprocess.run([str(soffice), "--headless", f"-env:UserInstallation={Path(profile).as_uri()}",
                            "--convert-to", "pdf", "--outdir", str(output), str(pptx_path)],
                           check=True, timeout=180)
            renderer_version = subprocess.check_output([str(soffice), "--headless",
                f"-env:UserInstallation={Path(profile).as_uri()}", "--version"], text=True, timeout=60).strip()
        renderer = "LibreOffice native PPTX-to-PDF export"
    assets = output / "slide_assets"
    assets.mkdir(exist_ok=True)
    document = fitz.open(pdf_path)
    assert len(document) == 16
    contact = Image.new("RGB", (1920, 1080), "white")
    for index, page in enumerate(document):
        assert page.get_text().strip(), f"PDF slide {index + 1} has no text"
        pixmap = page.get_pixmap(matrix=fitz.Matrix(.5, .5), alpha=False)
        tile = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        tile.thumbnail((480, 270))
        contact.paste(tile, ((index % 4) * 480, (index // 4) * 270))
        if index in [2, 5, 7, 8, 10, 13]:
            page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(str(assets / f"slide_{index+1:02}.png"))
    contact.save(assets / "contact_sheet.jpg", quality=85)
    write_json(output / "slide_manifest.json", {
        "slides": slides, "team": team, "speaker_map": speaker_map, "main_slides": 12, "backup_slides": 4,
        "rendered_pdf_pages": len(document), "renderer": renderer, "renderer_version": renderer_version,
        "talk_seconds": 720, "qa_seconds": 180,
        "role_seconds": {role: 180 for role in ["Part 1", "Part 2", "Part 3", "Part 4"]},
        "editable": "Native PowerPoint text, shapes and charts with embedded workbooks",
        "source_csv_sha256": file_hash(ROOT / "results/final/model_comparison.csv"),
        "interpretation_sources": {name: file_hash(ROOT / "results/rf" / name) for name in
            ["validation_importance.csv", "misclassification_cases_05.csv", "test_predictions.csv"]},
        "generator_sha256": file_hash(Path(__file__)),
        "artifacts": {p.name: file_hash(p) for p in [pptx_path, pdf_path, notes_path]}})
    print(f"Created {len(slides)} editable slides, PDF, English notes and manifest.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--soffice", type=Path, help="Optional installed or portable LibreOffice executable")
    run(parser.parse_args().soffice)
