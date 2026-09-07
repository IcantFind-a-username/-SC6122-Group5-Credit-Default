"""Render the personal study guide to PDF/DOCX; no model fitting or selection."""
from html import escape
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import fitz

from integration.artifacts import file_hash, write_json

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'
STEM = 'Group5_Part4_Understanding_Guide'
FONT_DIR = Path('/System/Library/Fonts')


def blocks(text):
    """Parse the small heading/list/table subset used by this authored guide."""
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            yield 'break', ''
        elif line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip())
                index += 1
            yield 'table', [[cell.strip() for cell in row.strip('|').split('|')]
                            for row in rows if not re.fullmatch(r'[|\s:\-]+', row)]
        elif line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            yield 'h' + str(level), line[level:].strip()
        elif line.startswith('- '):
            yield 'bullet', line[2:]
        else:
            yield 'p', line


def inline(text):
    value = escape(text)
    value = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', value)
    return re.sub(r'`([^`]+)`', r'<code>\1</code>', value)


def add_runs(paragraph, text):
    for index, segment in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        run = paragraph.add_run(segment.replace('`', ''))
        run.bold = index % 2 == 1


def run():
    source = OUT / (STEM + '.md')
    parsed = list(blocks(source.read_text()))
    docx = Document()
    section = docx.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    for name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
        style = docx.styles[name]
        style.font.name = 'Arial'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Heiti SC')
    normal = docx.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        docx.styles[name].font.color.rgb = RGBColor.from_string('183447')
    footer = section.footer.paragraphs[0]
    footer.text = 'SC6122 Group 5 · Xu Yiqun · Personal study guide   '
    page_field = OxmlElement('w:fldSimple')
    page_field.set(qn('w:instr'), 'PAGE')
    footer._p.append(page_field)
    chapters = [[]]
    for kind, item in parsed:
        if kind == 'break':
            chapters.append([])
            docx.add_page_break()
        elif kind == 'table':
            table = docx.add_table(rows=0, cols=len(item[0]))
            table.style = 'Light Shading Accent 1'
            for row_index, row in enumerate(item):
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    add_runs(cell.paragraphs[0], text)
                    for run_value in cell.paragraphs[0].runs:
                        run_value.font.size = Pt(10)
                        if row_index == 0:
                            run_value.bold = True
            chapters[-1].append('<table>' + ''.join(
                '<tr>' + ''.join(f'<{"th" if index == 0 else "td"}>{inline(cell)}</{"th" if index == 0 else "td"}>'
                                 for cell in row) + '</tr>' for index, row in enumerate(item)) + '</table>')
        elif kind.startswith('h'):
            add_runs(docx.add_heading(level=min(int(kind[1:]), 3)), item)
            chapters[-1].append(f'<{kind}>{inline(item)}</{kind}>')
        else:
            paragraph = docx.add_paragraph(style='List Bullet' if kind == 'bullet' else 'Normal')
            add_runs(paragraph, item)
            chapters[-1].append('<p>' + ('• ' if kind == 'bullet' else '') + inline(item) + '</p>')
    docx.save(OUT / (STEM + '.docx'))
    css = '''
    @font-face {font-family: Study; src: url("STHeiti Light.ttc");}
    @font-face {font-family: Study; font-weight: bold; src: url("STHeiti Medium.ttc");}
    body {font-family: Study; font-size:11pt; line-height:1.48; color:#233640;}
    h1 {font-size:22pt; color:#183447; margin:0 0 12pt;}
    h2 {font-size:18pt; color:#183447; margin:0 0 12pt;}
    h3 {font-size:12.5pt; color:#187B78; margin:12pt 0 5pt;}
    p {margin:0 0 8pt;}
    b {font-weight:bold;}
    table {width:100%; border-collapse:collapse; font-size:9.5pt; margin:8pt 0 12pt;}
    th,td {padding:5pt; border-bottom:0.5pt solid #d5dfe1; vertical-align:top;}
    th {background:#edf3f3; text-align:left;}
    code {font-family: Study;}
    '''
    raw = OUT / (STEM + '.render.pdf')
    writer = fitz.DocumentWriter(str(raw))
    bounds = fitz.Rect(0, 0, 595, 842)
    content = fitz.Rect(45, 43, 550, 789)
    page_count = 0
    for chapter in chapters:
        story = fitz.Story(html='<html><body>' + '\n'.join(chapter) + '</body></html>',
                           user_css=css, archive=fitz.Archive(str(FONT_DIR)))
        more = True
        while more:
            page_count += 1
            if page_count > 30:
                raise ValueError('Unexpected guide overflow')
            device = writer.begin_page(bounds)
            more, _ = story.place(content)
            story.draw(device)
            writer.end_page()
    writer.close()
    document = fitz.open(raw)
    for index, page in enumerate(document):
        page.insert_text((45, 818), 'SC6122 Group 5 | Xu Yiqun | Personal study guide', fontsize=8, color=(.35,.43,.46))
        page.insert_text((533, 818), str(index+1), fontsize=9, color=(.35,.43,.46))
    document.subset_fonts()
    pdf = OUT / (STEM + '.pdf')
    document.save(pdf, garbage=4, deflate=True)
    document.close()
    raw.unlink()
    result = fitz.open(pdf)
    all_text = '\n'.join(page.get_text() for page in result)
    for expected in ['Xu Yiqun', 'G2509092H', '18,000', '0.5598', '0.5620', '7.52%', '3,408']:
        assert expected in all_text, expected
    assert '阈值' in all_text and '成本' in all_text
    write_json(OUT / 'study_guide_manifest.json', {
        'pdf_pages': len(result), 'purpose': 'Personal learning guide, not the final group report',
        'source_SHA256': file_hash(source),
        'files': {file.name: file_hash(file) for file in [source, pdf, OUT/(STEM+'.docx')]},
        'model_comparison_SHA256': file_hash(ROOT/'results/final/model_comparison.csv'),
        'protocol_SHA256': file_hash(ROOT/'results/xgboost/protocol_frozen.json'),
        'cv_results_SHA256': file_hash(ROOT/'results/xgboost/cv_results.csv')})
    print(f'Generated {len(result)}-page PDF, editable DOCX and source Markdown.')


if __name__ == '__main__':
    run()
