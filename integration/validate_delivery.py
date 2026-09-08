"""Render and mechanically check delivered documents against accepted metrics."""
import argparse
import json
import math
from pathlib import Path
import re
import unicodedata

import fitz
from PIL import Image
from pptx import Presentation

from integration.artifacts import file_hash, write_json
from integration.audit import read_csv
from integration.presentation_evidence import rf_interpretation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'


def run(report_only=False):
    rendered = OUT/'rendered'
    rendered.mkdir(exist_ok=True)
    pdf = fitz.open(OUT/'Group5_Final_Report.pdf')
    print('Report page openings:', [(i+1, p.get_text()[:100]) for i,p in enumerate(pdf)])
    report_text = '\n'.join(p.get_text() for p in pdf)
    supplied_style = ROOT/'course_materials/latex_template/neurips_2021.sty'
    assert file_hash(OUT/'neurips_2021.sty') == file_hash(supplied_style)
    assert all(abs(page.rect.width-612) < .1 and abs(page.rect.height-792) < .1 for page in pdf)
    assert 'Abstract' in report_text and 'AI disclosure' in report_text
    assert 'language polishing and grammar guidance' in report_text
    assert 'Anonymous Author' not in report_text and 'Under review' not in report_text
    assert 'Conference on Neural Information Processing Systems' not in report_text
    fonts = {font[0]: font for page in pdf for font in page.get_fonts(full=True)}
    assert all(font[2] != 'Type3' and pdf.extract_font(xref)[3] for xref,font in fonts.items())
    assert 'Overfull' not in (OUT/'Group5_Final_Report.log').read_text()
    team = json.loads((OUT/'team.json').read_text())
    qa = (OUT/'Group5_QA.md').read_text()
    assert sum(member['share_percent'] for member in team) == 100
    for member in team:
        assert member['name'] in report_text and member['student_id'] in report_text
        assert member['name'] in qa and member['student_id'] in qa
        assert member['share_percent'] == 25
    assert not re.search(r'TO CONFIRM|\bTBD\b|NAME /', report_text)
    original_lr = read_csv(ROOT/'results/logistic_regression_final_comparison.csv')
    original_lr = original_lr[original_lr.Model == 'Baseline LR'].iloc[0]
    assert math.isclose(original_lr.Accuracy, (original_lr.TN+original_lr.TP)/6000)
    assert math.isclose(original_lr.Recall, original_lr.TP/(original_lr.TP+original_lr.FN))
    for key in ['Accuracy', 'ROC-AUC', 'Recall']:
        assert f'{original_lr[key]:.4f}' in report_text
    rf_importance, rf_examples = rf_interpretation()
    assert f'{rf_importance.iloc[0].Mean_AP_decrease:.4f}' in report_text
    for _, example in rf_examples.iterrows():
        assert str(int(example.row_id)) in report_text
        assert f'{example.Tuned_probability:.4f}' in report_text
    for text in ['Problem','References','contribution','hypothetical','post-hoc']:
        assert text.lower() in report_text.lower(), text
    test = read_csv(ROOT/'results/final/model_comparison.csv')
    test = test[(test.Partition=='test') & (test.Threshold==.5)]
    for _, row in test.iterrows():
        for key in ['AP','ROC-AUC','Precision','Recall','F1','Accuracy']:
            assert f'{row[key]:.4f}' in report_text, (row.Model,key)
    pages = []
    for i,page in enumerate(pdf):
        spans=[s for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for s in line['spans']]
        # Three-point tolerance permits normal glyph protrusion in the 5.5-inch text block.
        outside=[s['text'] for s in spans if s['bbox'][0]<105 or s['bbox'][2]>507 or s['bbox'][3]>page.rect.height-20]
        assert not outside, (i+1,outside)
        page.get_pixmap(matrix=fitz.Matrix(1,1)).save(rendered/f'report_page_{i+1}.png')
        pages.append({'page':i+1,'words':len(page.get_text().split()),'min_font_pt':round(min(s['size'] for s in spans),2),
                      'rightmost_text_pt':round(max(s['bbox'][2] for s in spans),2)})
    sheet=Image.new('RGB',(1290,((len(pdf)+2)//3)*610),'#dae0e6')
    for i in range(len(pdf)):
        im=Image.open(rendered/f'report_page_{i+1}.png').convert('RGB')
        im.thumbnail((420,594))
        sheet.paste(im,((i%3)*430,(i//3)*610))
    sheet.save(rendered/'report_contact.png')
    result={'report_pages':len(pdf),'report_page_details':pages,
            'teacher_template': 'Supplied neurips_2021.sty is byte-identical; US Letter; native 10pt typography; named-author mode; course-only notice',
            'teacher_style_SHA256':file_hash(supplied_style),
            'fonts_embedded_no_type3':True, 'overfull_tex_boxes':0,
            'report_comparison_numbers':'All baseline/tuned 0.5 numeric metrics found in PDF',
            'report_pdf_SHA256':file_hash(OUT/'Group5_Final_Report.pdf'), 'report_latex_SHA256':file_hash(OUT/'Group5_Final_Report.tex'),
            'visual_inspection':'Rendered contact sheet and full pages are separately inspected; automated bounds do not alone establish visual quality.'}
    if not report_only and (OUT/'Group5_Presentation.pdf').exists():
        deck=Presentation(str(OUT/'Group5_Presentation.pptx'))
        slides_pdf=fitz.open(OUT/'Group5_Presentation.pdf')
        assert len(deck.slides)==len(slides_pdf)==16
        manifest = json.loads((OUT/'slide_manifest.json').read_text())
        assert manifest['team'] == team
        for name, digest in manifest['interpretation_sources'].items():
            assert digest == file_hash(ROOT/'results/rf'/name)
        assert manifest['talk_seconds'] == 720 and manifest['qa_seconds'] == 180
        assert manifest['role_seconds'] == {f'Part {i}':180 for i in range(1,5)}
        assert manifest['source_csv_sha256'] == file_hash(ROOT/'results/final/model_comparison.csv')
        notes = (OUT/'Group5_Speaker_Notes.md').read_text()
        def normalized(text):
            return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKC',text).lower())
        text_checks = 0
        for number,(slide,page) in enumerate(zip(deck.slides,slides_pdf),1):
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    assert normalized(shape.text) in normalized(page.get_text()), (number,shape.text)
                    text_checks += 1
            spec = manifest['slides'][number-1]
            assert spec['script'] in notes and spec['transition'] in notes
            native = slide.notes_slide.notes_text_frame.text
            assert spec['transition'] in native and spec['script'] in native
            speaker = manifest['speaker_map'][str(number)]
            member = next(m for m in team if m['name'] == speaker)
            assert member['name'] in native and member['student_id'] in native
            assert member['name'] in page.get_text()
            assert not re.search(r'TO CONFIRM|\bTBD\b|NAME /|\bMIN\b|\bSEC\b', page.get_text())
        for member in team:
            assert member['name'] in slides_pdf[0].get_text()
            assert member['student_id'] in slides_pdf[0].get_text()
        rf_chart = next(shape.chart for shape in deck.slides[7].shapes if shape.has_chart)
        expected = rf_importance.head(5).iloc[::-1]
        assert [category.label for category in rf_chart.plots[0].categories] == expected.Feature.tolist()
        assert all(math.isclose(actual, value, abs_tol=1e-12) for actual, value in
                   zip(rf_chart.series[0].values, expected.Mean_AP_decrease, strict=True))
        for _, example in rf_examples.iterrows():
            assert str(int(example.row_id)) in slides_pdf[8].get_text()
            assert f'{example.Tuned_probability:.4f}' in slides_pdf[8].get_text()
        result['rf_interpretation'] = 'Native importance chart matches frozen validation data; illustrated error scores/IDs match saved predictions and both final documents'
        result['member_identity_and_notes'] = 'Four named members; matching student IDs; 25% each; all 16 slides have complete native scripts and named leads'
        result['slide_text_boxes_verified'] = text_checks
        result['speaker_scripts_and_transitions'] = 'Match manifest, Markdown and native speaker notes'
        result.update({'slide_pages':len(slides_pdf),'slide_first_text_matches_pdf':True,
                       'presentation_pptx_SHA256':file_hash(OUT/'Group5_Presentation.pptx'),
                       'presentation_pdf_SHA256':file_hash(OUT/'Group5_Presentation.pdf')})
    assert len(pdf) <= 7, f'Report has {len(pdf)} pages'
    write_json(OUT/('report_validation.json' if report_only else 'delivery_validation.json'),result)
    print(json.dumps(result,indent=2))


if __name__=='__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report-only', action='store_true', help='Validate the report without touching a separately edited presentation.')
    run(report_only=parser.parse_args().report_only)
